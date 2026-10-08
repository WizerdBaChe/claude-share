# -*- coding: utf-8 -*-
"""deck_layout -- layout layer over deck_builder.py for AUDIENCE decks (presentation class).

Ported 2026-09-21 from COMSOL_Test round 32 r2 (32_tutorial-decks/deck_kit.py), whose v1 was rejected at
the design layer for missing type hierarchy, a title/description merge and text-only example pages. It
turns rules/office-deck-deliverables.md P1-P5 into CONSTANTS and helpers instead of habits:

  * P5 type tiers: title 28 / description 16 (its own text object, head()) / subheading 19 / body 15.5
    (floor 14) / table 13-15 / caption 12 / kicker + page number 11.5; body at 1.5 line spacing.
    block() searches size only inside the tier and REPORTS what does not fit (FINDINGS), never shrinks
    below the floor.
  * P1: declare(slide, prs, kind, question) writes the page's figure kind + question into the notes;
    a build lists undeclared pages.
  * Figure card for beginners: figcard(kind "zh (en)", axes, goal) under each figure; fig_row() lays
    figures + caption + card side by side.
  * Height estimate calibrated on a PowerPoint COM render: with a line-spacing multiple the pitch is
    1.3 x size x spacing (single spacing keeps 1.55).
  * Orphan wraps (<= 8 chars on the last line): first a soft break after a clause mark (clause_break,
    both halves >= 0.3 w, never before punctuation), then narrow the box up to 20 %, then report.
    The tail estimate uses RAW width; budgets use the inflated width.
  * P3 flat: rectangles get an explicit fill + shadow.inherit = False; connectors get an empty
    a:effectLst so the theme's effects are not inherited.
  * type_audit(pptx): smallest run per text object read back from the EMITTED file.

Visual acceptance is still the COM render read page by page; these gates only shrink what it must catch.
Run `python deck_layout.py` for a 3-page sample (layout_sample.pptx) + the type audit.
"""
from __future__ import annotations

import math
import re
from pathlib import Path

from PIL import Image
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.util import Emu, Inches, Pt

import deck_builder as H

W, HH = 13.333, 7.5
MX = 0.6
CW = W - 2 * MX
BOT = HH - 0.5            # page number band starts here

# ---- P5 type tiers (pt) ----
KICKER_PT = 11.5
TITLE_PT = 28
DESC_PT = 16
SUB_PT = 19
BODY_PT = 15.5
BODY_FLOOR = 14
TABLE_PT, TABLE_FLOOR, TABLE_CEIL = 14, 13, 15
CAP_PT = 12
PN_PT = 11.5
LINE_SP = 1.5
PARA_GAP_PT = 6

INK = H.INK
MUTED = H.MUTED
ACCENT = H.GLASS                      # the one accent colour
WARN = H.FAIL_C                       # used only where a page contrasts right / wrong
SOFT = RGBColor(0xF4, 0xF6, 0xF8)
HEAD_BG = RGBColor(0x2B, 0x38, 0x47)
ZEBRA = RGBColor(0xF6, 0xF8, 0xFA)
CODE_BG = RGBColor(0xF1, 0xF3, 0xF5)

FINDINGS: list[dict] = []
DECL: list[dict] = []


def new_prs():
    prs = Presentation()
    prs.slide_width = Inches(W)
    prs.slide_height = Inches(HH)
    return prs


def blank(prs):
    return H.blank(prs)


def notes(s, text):
    H.notes(s, text)


# ------------------------------------------------------------------ declarations (P1)
def declare(s, prs, kind, question):
    """kind: 原生圖 / 數據圖 / 表 / 示意圖 / 流程圖 / 文字（附理由）"""
    DECL.append({"page": len(prs.slides._sldIdLst), "kind": kind, "question": question})


# ------------------------------------------------------------------ text with height discipline
def _segs(p):
    if isinstance(p, str):
        return H_md(p)
    if isinstance(p, dict):
        return _segs(p["segs"])
    return [(t, dict(o)) for t, o in p]


def H_md(text):
    s = str(text)
    parts = s.split("**")
    if len(parts) % 2 == 0:
        return [(s, {})]
    out = [(p, {"bold": True} if i % 2 else {}) for i, p in enumerate(parts) if p]
    return out or [("", {})]


def plain(text):
    return str(text).replace("**", "")


def est_h(paras, size, w, pitch=1.55, ls=LINE_SP):
    h = 0.0
    for p in paras:
        segs = _segs(p)
        sa = p.get("space_after", PARA_GAP_PT) if isinstance(p, dict) else PARA_GAP_PT
        sz = p.get("size", size) if isinstance(p, dict) else size
        txt = "".join(t for t, _ in segs)
        lines = 0
        for part in re.split("[\n\v]", txt):
            lines += max(1, math.ceil(H._est_w_in([(part, {})], sz) * 1.04 / w))
        pls = p.get("ls", ls) if isinstance(p, dict) else ls
        # Measured on this deck's COM render (r2 run 2, page 14 figure card): with a line-spacing multiple the real
        # pitch is ~1.2 x size x ls; 1.3 keeps the bias tall. The 1.55 pitch (L-078) stays for single spacing.
        h += lines * sz * (1.3 * pls if pls > 1.05 else pitch) / 72 + sa / 72
    return h


def orphan_tail(text, size, w):
    """Estimated width (in chars) of the last wrapped line; None when the text fits one line."""
    wt = H._est_w_in([(text, {})], size)      # RAW width for a tail (rule: inflation is for budgets only)
    if wt <= w:
        return None
    tail = wt - math.floor(wt / w) * w
    return tail / (size / 72)


BREAK_AFTER = "，；：、）。"
ORPHAN_MAX = 8.5          # chars; a last line shorter than this is an orphan (user ruling 2026-08-31: <= 8)


def clause_break(segs, size, w):
    """A paragraph that wraps onto a short last line and is at most two lines long gets a soft break after the
    clause mark that best balances the two lines (both must fit). Returns new segs, or None when no clean split
    exists. The text is unchanged; only a line break is added."""
    txt = "".join(tt for tt, _ in segs)
    if "\v" in txt or "\n" in txt:
        return None
    wt = H._est_w_in([(txt, {})], size)
    if wt <= w or wt > 1.94 * w:
        return None
    best = None
    for i, ch in enumerate(txt[:-1]):
        if ch not in BREAK_AFTER or txt[i + 1] in BREAK_AFTER:   # never start a line with punctuation
            continue
        left = H._est_w_in([(txt[:i + 1], {})], size) * 1.04
        right = H._est_w_in([(txt[i + 1:], {})], size) * 1.04
        if left <= 0.97 * w and right <= 0.97 * w and min(left, right) >= 0.3 * w:
            score = abs(left - right)
            if best is None or score < best[0]:
                best = (score, i + 1)
    if best is None:
        return None
    idx, out, pos = best[1], [], 0
    for tt, o in segs:
        if pos < idx <= pos + len(tt):
            cut = idx - pos
            tt = tt[:cut] + "\v" + tt[cut:]
        out.append((tt, o))
        pos += len(tt.replace("\v", ""))
    return out


def block(s, x, y, w, h, paras, size=BODY_PT, floor=BODY_FLOOR, name="content", color=None,
          anchor=MSO_ANCHOR.TOP, align=None):
    """Paragraph list in ONE text object. A paragraph may be a dict with segs/size/bold/color/space_after/ls."""
    sz = size
    while est_h(paras, sz, w) > h and sz - 0.5 >= floor:
        sz -= 0.5
    need = est_h(paras, sz, w)
    if need > h + 0.02:
        FINDINGS.append({"gate": "fit", "shape": name, "need_in": round(need, 2), "box_in": round(h, 2),
                         "text": "".join(t for t, _ in _segs(paras[0]))[:30]})
    def orphans(ww):
        out = []
        for p in paras:
            o = p if isinstance(p, dict) else {}
            txt = "".join(tt for tt, _ in _segs(p))
            for part in re.split("[\n\v]", txt):
                tail = orphan_tail(part, o.get("size", sz), ww - 0.05)
                # lower bound 0: a near-fit is a risk too (r2 run-7 render: a 1-char tail the raw model put at <1)
                if tail is not None and tail < ORPHAN_MAX:
                    out.append({"gate": "orphan", "shape": name, "tail_chars": round(tail, 1), "text": part[-14:]})
        return out
    # First remedy: a soft break after a clause mark (two-line paragraphs only).
    fixed = []
    for p in paras:
        o = dict(p) if isinstance(p, dict) else {}
        segs = _segs(p)
        txt = "".join(tt for tt, _ in segs)
        tail = orphan_tail(txt, o.get("size", sz), w - 0.05) if "\n" not in txt else None
        if tail is not None and tail < ORPHAN_MAX:
            nb = clause_break(segs, o.get("size", sz), w - 0.05)
            if nb:
                o["segs"] = nb
                fixed.append(o)
                continue
        fixed.append(p)
    paras = fixed
    # Balance before reporting: narrowing the box pushes words onto the short last line (never widens past the
    # layout's box, never changes the text). Accept a width only if the estimated height still fits.
    bad = orphans(w)
    if bad:
        for k in range(1, 21):
            ww = w * (1 - 0.01 * k)
            if not orphans(ww) and est_h(paras, sz, ww) <= max(h, need) + 0.02:
                w, bad = ww, []
                break
    FINDINGS.extend(bad)
    box, tf = H.tb(s, x, y, w, h)
    box.name = name
    tf.vertical_anchor = anchor
    shrink = sz - size
    for i, p in enumerate(paras):
        segs = _segs(p)
        o = p if isinstance(p, dict) else {}
        psz = o.get("size", size) + (shrink if "size" not in o else 0)
        pp = H.para(tf, segs, size=psz, color=o.get("color", color or INK), bold=o.get("bold", False),
                    first=(i == 0), space_after=o.get("space_after", PARA_GAP_PT), align=o.get("align", align),
                    line_spacing=o.get("ls", LINE_SP), name=o.get("name", H.SANS))
        if o.get("hang"):   # hanging indent in characters: wrapped lines align under the text, not the label
            ppr = pp._p.get_or_add_pPr()
            emu = int(o["hang"] * psz * 12700)
            ppr.set("marL", str(emu))
            ppr.set("indent", str(-emu))
    return sz


# ------------------------------------------------------------------ page head (P5: title and description apart)
def head(s, kicker, title, desc=None):
    """kicker (small accent label) / title 28 pt / description 16 pt as a SEPARATE object. Returns content top y."""
    y = 0.34
    if kicker:
        box, tf = H.tb(s, MX, y, CW, 0.3)
        box.name = "kicker"
        H.para(tf, kicker, size=KICKER_PT, color=ACCENT, bold=True, first=True, space_after=0)
        y += 0.32
    th = 0.62
    box, tf = H.tb(s, MX, y, CW, th)
    box.name = "title"
    H.para(tf, H_md(title), size=TITLE_PT, bold=True, first=True, space_after=0, fit_w=CW)
    y += th + 0.08
    if desc:
        dh = max(0.42, est_h([desc], DESC_PT, CW, ls=1.3))
        block(s, MX, y, CW, dh, [{"segs": H_md(desc), "ls": 1.3, "space_after": 0}], size=DESC_PT,
              floor=DESC_PT - 1, name="description", color=MUTED)
        y += dh
    return y + 0.28


# ------------------------------------------------------------------ figures
def fig(s, path, x, y, w, h, align="center"):
    """Fit an image into (w, h) keeping its aspect (measured, never guessed). Returns the drawn box."""
    path = Path(path)
    im = Image.open(path)
    ar = im.height / im.width
    fw, fh = (w, w * ar) if w * ar <= h else (h / ar, h)
    fx = x + (w - fw) / 2 if align == "center" else x
    pic = H.add_pic(s, path, fx, y, w=fw)
    pic.name = "figure"
    return fx, y, fw, fh


def caption(s, x, y, w, text, h=0.34, align=None):
    block(s, x, y, w, h, [{"segs": H_md(text), "ls": 1.15, "space_after": 0}], size=CAP_PT, floor=CAP_PT - 1,
          name="caption", color=MUTED, align=align)


def point(s, x, y, w, h, sub, sentence, name="point", sub_color=None):
    """Subheading (19 pt bold) + one sentence (15.5 pt) in one text object."""
    paras = [{"segs": H_md(sub), "size": SUB_PT, "bold": True, "color": sub_color or INK, "ls": 1.15,
              "space_after": 4}]
    if sentence:
        paras.append({"segs": H_md(sentence), "space_after": 0})
    return block(s, x, y, w, h, paras, name=name)


def point_h(sub, sentence, w):
    paras = [{"segs": H_md(sub), "size": SUB_PT, "ls": 1.15, "space_after": 4}]
    if sentence:
        paras.append({"segs": H_md(sentence), "space_after": 0})
    return est_h(paras, BODY_PT, w)


def points(s, x, y, w, h, items, name="points", gap_min=0.25):
    """Stack (subheading, sentence) blocks and spread them over the free height."""
    hs = [point_h(a, b, w) + 0.04 for a, b in items]
    free = h - sum(hs)
    gap = max(gap_min, free / max(1, len(items) - 1)) if len(items) > 1 else 0
    if sum(hs) + gap_min * (len(items) - 1) > h + 0.02:
        FINDINGS.append({"gate": "fit", "shape": name, "need_in": round(sum(hs) + gap_min * (len(items) - 1), 2),
                         "box_in": round(h, 2), "text": items[0][0]})
        gap = gap_min
    gap = min(gap, 0.6)
    yy = y
    for i, (a, b) in enumerate(items):
        point(s, x, yy, w, hs[i], a, b, name=f"{name}{i + 1}")
        yy += hs[i] + gap
    return yy


FIGKIND_PT = 18
FIGCARD_PT = 15


def figcard_paras(kind, axes, goal):
    zh, _, en = kind.partition(" (")
    segs = [(zh, {"bold": True})] + ([("  " + en.rstrip(")"), {"bold": False, "size": 14, "color": MUTED})] if en else [])
    paras = [{"segs": segs, "size": FIGKIND_PT, "color": ACCENT, "ls": 1.15, "space_after": 4}]
    if axes:
        paras.append({"segs": [("軸　", {"bold": True, "color": MUTED})] + H_md(axes), "size": FIGCARD_PT, "space_after": 3, "hang": 2})
    if goal:
        paras.append({"segs": [("看　", {"bold": True, "color": MUTED})] + H_md(goal), "size": FIGCARD_PT, "space_after": 0, "hang": 2})
    return paras


def figcard(s, x, y, w, h, kind, axes, goal, name="figcard"):
    """The chart explained for a beginner: what kind of chart (zh + en), what the axes are, what to read off it."""
    return block(s, x, y, w, h, figcard_paras(kind, axes, goal), size=FIGCARD_PT, floor=14, name=name)


def figcard_h(kind, axes, goal, w):
    return est_h(figcard_paras(kind, axes, goal), FIGCARD_PT, w)


def fig_row(s, y, items, fig_h, bottom=None, gap=0.35, x0=MX, total=CW):
    """items: [dict(path, cap, kind, axes, goal, share=1)] -> figures in one row, caption + figure card under each.
    Returns y below the tallest card."""
    bottom = bottom or BOT - 0.05
    shares = [it.get("share", 1) for it in items]
    ws = [(total - gap * (len(items) - 1)) * sh / sum(shares) for sh in shares]
    x = x0
    ymax = y
    for it, w in zip(items, ws):
        fx, fy, fw, fh = fig(s, it["path"], x, y, w, fig_h)
        cy = y + fig_h + 0.04
        # two-line budget: a reader's one-line addition must not reach the card (headroom COLLIDE,
        # measured 16.7 pt on the open-slide-borrow pilot deck, 2026-09-29)
        caption(s, x, cy, w, it["cap"], h=0.55)
        ky = cy + 0.6
        if it.get("kind"):
            figcard(s, x, ky, w, bottom - ky, it["kind"], it.get("axes"), it.get("goal"), name="figcard")
            ymax = max(ymax, ky + figcard_h(it["kind"], it.get("axes"), it.get("goal"), w))
        x += w + gap
    return ymax


def points_row(s, x, y, w, h, items, gap=0.45, name="points"):
    """(subheading, sentence) blocks side by side: the flat carrier for 2-3 parallel points under a table."""
    cw = (w - gap * (len(items) - 1)) / len(items)
    for i, (a, b) in enumerate(items):
        point(s, x + i * (cw + gap), y, cw, h, a, b, name=f"{name}{i + 1}")


def rule(s, x, y, w, color=None, t=0.035):
    return H.add_rect(s, x, y, w, t, color or ACCENT)


# ------------------------------------------------------------------ tables (T16 sizing, P5 tier)
ROW_FLOOR_IN = 0.36


def _row_h(r, sz, cw, infl=1.04):
    need, lines = 0.0, 1
    for c, v in enumerate(r):
        avail = max(cw[c] - 0.16, 0.3)
        n = 0
        parts = plain(v).split("\n")
        for part in parts:
            n += max(1, math.ceil(H._est_w_in([(part, {})], sz) * infl / avail))
        lines = max(lines, n)
        need = max(need, n * sz * 1.42 * 1.25 / 72 + (len(parts) - 1) * PARA_GAP_PT / 72)
    return max(need + 0.1, lines * ROW_FLOOR_IN)


def table(s, x, y, w, header, rows, col_w=None, max_h=None, name="table", first_bold=True, hl_rows=()):
    col_w = col_w or [1] * len(header or rows[0])
    tot = sum(col_w)
    cw = [w * c / tot for c in col_w]
    allrows = ([header] if header else []) + rows
    sz = TABLE_CEIL
    while sz > TABLE_FLOOR and max_h and sum(_row_h(r, sz, cw) for r in allrows) > max_h:
        sz -= 0.5
    budget = sum(_row_h(r, sz, cw) for r in allrows)
    if max_h and budget > max_h + 0.02:
        FINDINGS.append({"gate": "fit", "shape": name, "need_in": round(budget, 2), "box_in": round(max_h, 2),
                         "text": f"table {len(rows)} rows"})
    fixed_rows = []
    for r in allrows:
        nr = []
        for c, v in enumerate(r):
            parts = []
            for part in str(v).split(chr(10)):
                tail = orphan_tail(plain(part), sz, cw[c] - 0.16)
                if tail is not None and tail < ORPHAN_MAX:
                    nb = clause_break([(part, {})], sz, cw[c] - 0.18) if "**" not in part else None
                    if nb:
                        part = nb[0][0]
                    else:
                        FINDINGS.append({"gate": "orphan", "shape": name, "tail_chars": round(tail, 1), "text": part[-14:]})
                parts.append(part)
            nr.append(chr(10).join(parts))
        fixed_rows.append(nr)
    allrows = fixed_rows
    hs = [_row_h(r, sz, cw, infl=1.01) for r in allrows]
    shp = s.shapes.add_table(len(allrows), len(cw), Inches(x), Inches(y), Inches(w), Inches(sum(hs)))
    shp.name = name
    tbl = shp.table
    tbl.horz_banding = False
    tbl.first_row = bool(header)
    for i, c in enumerate(cw):
        tbl.columns[i].width = Inches(c)
    for ri, r in enumerate(allrows):
        tbl.rows[ri].height = Inches(hs[ri])
        for ci, v in enumerate(r):
            cell = tbl.cell(ri, ci)
            cell.margin_left = cell.margin_right = Emu(70000)
            cell.margin_top = cell.margin_bottom = Emu(25000)
            cell.vertical_anchor = MSO_ANCHOR.MIDDLE
            cell.fill.solid()
            is_head = header and ri == 0
            dri = ri - (1 if header else 0)
            if is_head:
                cell.fill.fore_color.rgb = HEAD_BG
            elif dri in hl_rows:
                cell.fill.fore_color.rgb = H.GLASS_BG
            else:
                cell.fill.fore_color.rgb = ZEBRA if dri % 2 else H.WHITE
            tf = cell.text_frame
            tf.word_wrap = True
            parts = str(v).split("\n")
            for k, part in enumerate(parts):
                H.para(tf, H_md(part), size=sz, color=H.WHITE if is_head else INK,
                       bold=is_head or (first_bold and ci == 0), first=(k == 0),
                       space_after=PARA_GAP_PT if k < len(parts) - 1 else 0, line_spacing=1.25)
    return budget


# ------------------------------------------------------------------ diagrams
def boxflow(s, steps, y, h, x0=MX, total=CW, gap=0.4, fill=None, sub_pt=17, body_pt=14, name="flow"):
    """Left-to-right boxes with arrows. steps: [(label, body)]"""
    from pptx.enum.shapes import MSO_CONNECTOR
    from pptx.oxml.ns import qn
    n = len(steps)
    w = (total - gap * (n - 1)) / n
    for i, (k, body) in enumerate(steps):
        x = x0 + i * (w + gap)
        r = H.add_rect(s, x, y, w, h, fill or H.GLASS_BG, line_color=H.LINE)
        r.name = f"{name}-bg{i + 1}"
        paras = [{"segs": H_md(k), "size": sub_pt, "bold": True, "color": ACCENT, "ls": 1.1, "align": PP_ALIGN.CENTER,
                  "space_after": 4}]
        if body:
            paras.append({"segs": H_md(body), "size": body_pt, "ls": 1.25, "align": PP_ALIGN.CENTER, "space_after": 0})
        block(s, x + 0.12, y + 0.1, w - 0.24, h - 0.2, paras, size=body_pt, floor=13, name=f"{name}{i + 1}",
              anchor=MSO_ANCHOR.MIDDLE)
        if i < n - 1:
            c = s.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, Inches(x + w + 0.05), Inches(y + h / 2),
                                       Inches(x + w + gap - 0.05), Inches(y + h / 2))
            c.line.color.rgb = MUTED
            c.line.width = Pt(1.75)
            ln = c.line._get_or_add_ln()
            ln.append(ln.makeelement(qn("a:tailEnd"), {"type": "triangle", "w": "med", "len": "med"}))
            sp = c._element.spPr
            sp.append(sp.makeelement(qn("a:effectLst"), {}))   # P3: no inherited theme effect on connectors
    return w


def code(s, x, y, w, h, lines, size=13, name="code"):
    r = H.add_rect(s, x, y, w, h, CODE_BG, line_color=H.LINE)
    r.name = f"{name}-bg"
    paras = [{"segs": [(ln, {})], "name": H.MONO, "ls": 1.15, "space_after": 0} for ln in lines]
    block(s, x + 0.18, y + 0.14, w - 0.36, h - 0.28, paras, size=size, floor=12, name=name)


# ------------------------------------------------------------------ cover / section / page numbers
def cover(prs, title, sub, meta):
    s = blank(prs)
    rule(s, MX, 2.35, 1.2, t=0.06)
    block(s, MX, 2.6, CW, 1.1, [{"segs": H_md(title), "size": 38, "bold": True, "ls": 1.1, "space_after": 0}],
          size=38, floor=34, name="cover-title")
    block(s, MX, 3.75, CW, 0.9, [{"segs": H_md(sub), "size": 20, "ls": 1.25, "space_after": 0}], size=20, floor=18,
          name="cover-sub", color=MUTED)
    block(s, MX, 5.9, CW, 0.7, [{"segs": H_md(meta), "size": 13, "ls": 1.25, "space_after": 0}], size=13, floor=12,
          name="cover-meta", color=MUTED)
    return s


def section(prs, num, title, desc):
    s = blank(prs)
    block(s, MX, 2.25, 2.0, 1.45, [{"segs": [(num, {})], "size": 66, "bold": True, "color": ACCENT, "ls": 1.0,
                                  "space_after": 0}], size=66, floor=60, name="section-num")
    block(s, MX, 3.75, CW, 0.8, [{"segs": H_md(title), "size": 32, "bold": True, "ls": 1.1, "space_after": 0}],
          size=32, floor=28, name="section-title")
    block(s, MX, 4.6, CW * 0.8, 1.0, [{"segs": H_md(desc), "size": 17, "ls": 1.4, "space_after": 0}], size=17,
          floor=16, name="section-desc", color=MUTED)
    return s


def number_pass(prs):
    n = len(prs.slides._sldIdLst)
    for i, s in enumerate(prs.slides, 1):
        if i == 1:
            continue
        box, tf = H.tb(s, W - 1.4, HH - 0.42, 0.8, 0.3)
        box.name = "page-number"
        H.para(tf, f"{i} / {n}", size=PN_PT, color=MUTED, first=True, align=PP_ALIGN.RIGHT, space_after=0)


# ------------------------------------------------------------------ P5 check on the EMITTED file
def type_audit(path):
    """Smallest run per text object class, read back from the built pptx."""
    prs = Presentation(str(path))
    small = []
    for pi, s in enumerate(prs.slides, 1):
        for shp in s.shapes:
            frames = []
            if shp.has_text_frame:
                frames.append((shp.name, shp.text_frame))
            if shp.has_table:
                for r in shp.table.rows:
                    for c in r.cells:
                        frames.append((shp.name, c.text_frame))
            for nm, tf in frames:
                for p in tf.paragraphs:
                    for r in p.runs:
                        if not r.text.strip() or r.font.size is None:
                            continue
                        pt = r.font.size.pt
                        floor = {"caption": CAP_PT - 1, "kicker": KICKER_PT, "page-number": PN_PT, "cover-meta": CAP_PT}.get(nm.split("-")[0] if nm.startswith("page") else nm, None)
                        if nm == "page-number":
                            floor = PN_PT
                        if floor is None:
                            floor = TABLE_FLOOR if shp.has_table else (12 if nm.startswith("code") else BODY_FLOOR)
                        if pt + 1e-6 < floor:
                            small.append({"page": pi, "shape": nm, "pt": pt, "text": r.text[:20]})
    return small


# ------------------------------------------------------------------ sample (python deck_layout.py)
def _sample(out):
    from PIL import Image as _I, ImageDraw as _D
    img = Path(out).with_name("_layout_sample_fig.png")
    im = _I.new("RGB", (900, 540), "white")
    d = _D.Draw(im)
    d.line([(60, 480), (860, 480)], fill="black", width=3)
    d.line([(60, 480), (60, 40)], fill="black", width=3)
    d.line([(60, 470), (300, 400), (500, 150), (840, 90)], fill=(62, 124, 171), width=6)
    im.save(img)
    prs = new_prs()
    cover(prs, "版面層樣張", "標題與描述分開、字級分層、每張圖附圖卡", "deck_layout.py sample")
    s = blank(prs)
    y = head(s, "樣張", "一張圖＋圖卡", "圖卡告訴新手：這是什麼圖、軸代表什麼、要看哪裡。")
    declare(s, prs, "數據圖＋圖卡", "這條曲線要看什麼")
    fig_row(s, y, [dict(path=img, cap="數據圖：樣例曲線", kind="轉移特性曲線 (transfer curve)",
                        axes="橫軸輸入、縱軸輸出", goal="曲線開始上升的位置就是門檻"),
                   dict(path=img, cap="同一張圖，換一種讀法", kind="飽和曲線 (saturation curve)",
                        axes="橫軸長度、縱軸效率", goal="曲線變平的地方就是「夠長了」")], 2.6)
    s = blank(prs)
    y = head(s, "樣張", "表格也走字級", "表格 13–15 pt；儲存格孤行先插軟換行，再報告。")
    declare(s, prs, "表格", "三個做法各自的代價")
    table(s, MX, y, CW, ["做法", "自由度", "說明"],
          [["做法甲", "約 27 億", "記憶體放不下：這一欄刻意寫長一點來觸發子句斷行，看它怎麼處理"],
           ["做法乙", "2.9 M", "本機解完"]], [3.0, 2.5, 6.6])
    number_pass(prs)
    prs.save(out)
    return out


if __name__ == "__main__":
    import json as _json
    import sys as _sys
    _sys.stdout.reconfigure(encoding="utf-8")
    o = _sample(Path(__file__).with_name("layout_sample.pptx"))
    small = type_audit(o)
    print(f"built {o.name}; below-tier runs: {len(small)}; findings: {len(FINDINGS)}")
    print(_json.dumps(FINDINGS, ensure_ascii=False, indent=1))
    assert not small, small
    # clause_break controls: a two-line text with a short tail splits at its clause mark (known-true);
    # the same length with no clause mark cannot be split (known-false).
    txt = "上一段先寫到接近整行的長度，後半句再短短收尾"
    w0 = H._est_w_in([(txt, {})], BODY_PT) * 0.9
    hit = clause_break([(txt, {})], BODY_PT, w0)
    miss = clause_break([(txt.replace("，", "的"), {})], BODY_PT, w0)
    assert hit and "\v" in hit[0][0] and miss is None, (hit, miss)
    print("clause_break controls: split", repr(hit[0][0]), "/ unsplittable -> None")
    print("== LAYOUT SAMPLE OK ==")
