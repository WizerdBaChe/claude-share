# -*- coding: utf-8 -*-
"""Find wrapped lines that open with a closing punctuation mark, read from PowerPoint's own line breaks.

    python line_start.py <deck.pptx> [--json out.json]     exit 0 clean, 1 hits, 3 uncalibrated
    python line_start.py --selftest                         the two controls alone

A closing mark (，、。；：！？）」 …) at the start of a wrapped line means PowerPoint laid the run out with
Latin line breaking instead of East Asian kinsoku. Cause measured 2026-09-29 (open-slide-borrow pilot, one
frame, four variants): a run with no `lang` attribute gets Latin rules; lang="zh-TW" alone fixes it,
eaLnBrk / hangingPunct alone do not. An estimator cannot see this - only the layout engine knows where a
line breaks - so the gate asks PowerPoint (COM `TextRange.Lines`).

Calibration runs on every invocation: one frame of CJK text laid out at a sweep of widths, once with the
run language stripped (known-false, must produce >= 1 hit) and once tagged zh-TW (known-true, must produce
0). If either control misbehaves the deck's hits are reported with no ruling (exit 3).
"""
from __future__ import annotations

import argparse
import json
import sys
import tempfile
from pathlib import Path

CLOSERS = "，、。；：！？）」』》〉,.;:!?)"
CONTROL_TEXT = ("這份簡報本身就是整合測試：開案、取用、建置、檢查都走新流程，"
                "過程中找到兩個素材缺陷和一個掃描缺陷，都已修正。")


def _frames(shape, prefix=""):
    name = prefix + shape.Name
    if shape.Type == 6:  # msoGroup
        for k in range(1, shape.GroupItems.Count + 1):
            yield from _frames(shape.GroupItems(k), name + "/")
        return
    if shape.HasTextFrame and shape.TextFrame.HasText:
        yield name, shape.TextFrame.TextRange
    if shape.HasTable:
        t = shape.Table
        for r in range(1, t.Rows.Count + 1):
            for c in range(1, t.Columns.Count + 1):
                cs = t.Cell(r, c).Shape
                if cs.TextFrame.HasText:
                    yield f"{name}[{r},{c}]", cs.TextFrame.TextRange


def scan(app, path: Path) -> list[dict]:
    """Every wrapped line (not a paragraph's first line) whose first character is a closing mark."""
    pres = app.Presentations.Open(str(path), ReadOnly=True, Untitled=False, WithWindow=False)
    hits = []
    try:
        for i in range(1, pres.Slides.Count + 1):
            sl = pres.Slides(i)
            for k in range(1, sl.Shapes.Count + 1):
                for name, tr in _frames(sl.Shapes(k)):
                    starts = {tr.Paragraphs(p).Start for p in range(1, tr.Paragraphs().Count + 1)}
                    for j in range(2, tr.Lines().Count + 1):
                        ln = tr.Lines(j)
                        if ln.Start not in starts and ln.Text[:1] in CLOSERS:
                            hits.append({"slide": i, "frame": name, "line": j, "text": ln.Text.strip()[:24],
                                         "lang": ln.Characters(1, 1).LanguageID})
    finally:
        pres.Close()
    return hits


def _control_deck(out: Path, tagged: bool) -> Path:
    from pptx import Presentation
    from pptx.oxml.ns import qn
    from pptx.util import Emu, Inches, Pt

    prs = Presentation()
    prs.slide_width, prs.slide_height = Inches(13.333), Inches(7.5)
    s = prs.slides.add_slide(prs.slide_layouts[6])
    w = 2.0
    while w <= 5.0:  # a width sweep, so some break lands on a closing mark whatever the font metrics
        tb = s.shapes.add_textbox(Inches(0.5), Inches(0.5), Inches(w), Inches(3))
        tf = tb.text_frame
        tf.word_wrap = True
        tf.margin_left = tf.margin_right = Emu(0)
        r = tf.paragraphs[0].add_run()
        r.text = CONTROL_TEXT
        r.font.size = Pt(16)
        rPr = r._r.get_or_add_rPr()
        for tag in ("a:latin", "a:ea"):
            el = rPr.makeelement(qn(tag), {"typeface": "Microsoft JhengHei"})
            rPr.append(el)
        if tagged:
            rPr.set("lang", "zh-TW")
            rPr.set("altLang", "en-US")
        w = round(w + 0.07, 2)
    p = out / f"line_start_{'true' if tagged else 'false'}.pptx"
    prs.save(str(p))
    return p


def calibrate(app) -> tuple[bool, bool, int, int]:
    with tempfile.TemporaryDirectory() as d:
        f = scan(app, _control_deck(Path(d), tagged=False))
        t = scan(app, _control_deck(Path(d), tagged=True))
    return len(f) >= 1, len(t) == 0, len(f), len(t)


def _app():
    import pythoncom
    import win32com.client
    pythoncom.CoInitialize()
    try:
        win32com.client.GetActiveObject("PowerPoint.Application")
        running = True
    except Exception:
        running = False
    return win32com.client.Dispatch("PowerPoint.Application"), running


def main() -> int:
    sys.stdout.reconfigure(encoding="utf-8")
    ap = argparse.ArgumentParser(description=__doc__.split("\n")[0])
    ap.add_argument("deck", nargs="?", type=Path)
    ap.add_argument("--json", type=Path, help="write hits and calibration as JSON")
    ap.add_argument("--selftest", action="store_true", help="run the two controls only")
    a = ap.parse_args()
    if not a.selftest and not a.deck:
        ap.error("deck or --selftest required")

    app, running = _app()
    try:
        ok_f, ok_t, nf, nt = calibrate(app)
        print(f"[controls] known-false (no lang) {nf} hit(s) {'OK' if ok_f else 'FAIL'} · "
              f"known-true (zh-TW) {nt} hit(s) {'OK' if ok_t else 'FAIL'}")
        if a.selftest:
            return 0 if ok_f and ok_t else 1
        hits = scan(app, a.deck.resolve())
    finally:
        if not running:
            app.Quit()

    calibrated = ok_f and ok_t
    for h in hits:
        print(f"    slide {h['slide']} · {h['frame']} · line {h['line']} · lang {h['lang']}: {h['text']}")
    if a.json:
        a.json.write_text(json.dumps({"deck": str(a.deck), "calibrated": calibrated, "controls": [nf, nt],
                                      "hits": hits}, ensure_ascii=False, indent=1), encoding="utf-8")
    if not calibrated:
        print(f"[line-start] UNCALIBRATED; {len(hits)} hit(s) reported, no ruling")
        return 3
    print(f"[line-start] {len(hits)} wrapped line(s) open with a closing mark")
    return 1 if hits else 0


if __name__ == "__main__":
    sys.exit(main())
