"""pptx-review: read a reviewer's PowerPoint comments and the live PowerPoint cursor as JSON.

Two read-only instruments for the deck feedback loop (borrowed 2026-09-29 from
open-slide's `/apply-comments` + `/current-slide`, re-cut for PowerPoint):

  comments <deck.pptx>   every comment on every slide, each resolved to the shape it
                         most likely refers to, with the METHOD and confidence of that
                         resolution; unresolvable comments are reported, never guessed.
  cursor                 which deck / slide / shapes the user has open and selected in
                         the running PowerPoint (attaches only; never starts PowerPoint).
  selftest               two-sided controls on a generated fixture + a fake COM app.

Neither command writes to a deck. Edits belong in the deck's build source.
"""
from __future__ import annotations

import argparse
import hashlib
import json
import re
import sys
import zipfile
from pathlib import Path, PurePosixPath
from xml.etree import ElementTree as ET

EMU_PER_PT = 12700
NEAR_TOLERANCE_EMU = 457200  # 0.5 in: a pin dropped just outside a shape's box still counts, flagged "near"

NS = {
    "p": "http://schemas.openxmlformats.org/presentationml/2006/main",
    "a": "http://schemas.openxmlformats.org/drawingml/2006/main",
    "r": "http://schemas.openxmlformats.org/officeDocument/2006/relationships",
    "p188": "http://schemas.microsoft.com/office/powerpoint/2018/8/main",
    "a16": "http://schemas.microsoft.com/office/drawing/2014/main",
    "rel": "http://schemas.openxmlformats.org/package/2006/relationships",
}
REL_MODERN = "http://schemas.microsoft.com/office/2018/10/relationships/comments"
REL_LEGACY = "http://schemas.openxmlformats.org/officeDocument/2006/relationships/comments"
QUOTE_RE = re.compile(r"[「『“\"]([^」』”\"]{2,})[」』”\"]")


def _local(tag: str) -> str:
    return tag.rsplit("}", 1)[-1]


def _resolve(base: str, target: str) -> str:
    parts: list[str] = []
    for seg in (PurePosixPath(base).parent / target).parts:
        if seg == "..":
            parts.pop()
        elif seg != ".":
            parts.append(seg)
    return "/".join(parts)


def _rels(z: zipfile.ZipFile, part: str) -> list[tuple[str, str, str]]:
    p = PurePosixPath(part)
    rels_name = f"{p.parent}/_rels/{p.name}.rels"
    if rels_name not in z.namelist():
        return []
    root = ET.fromstring(z.read(rels_name))
    return [(r.get("Id"), r.get("Type"), _resolve(part, r.get("Target"))) for r in root]


def _text(el: ET.Element) -> str:
    paras = []
    for para in el.iter(f"{{{NS['a']}}}p"):
        paras.append("".join(t.text or "" for t in para.iter(f"{{{NS['a']}}}t")))
    return "\n".join(paras).strip()


def _shapes(slide_root: ET.Element) -> list[dict]:
    """Top-level shapes with absolute boxes; group children are listed with the group's box
    mapped through chOff/chExt so a pin on a grouped label still lands on the label."""
    out: list[dict] = []
    tree = slide_root.find("p:cSld/p:spTree", NS)

    def walk(container, xform):
        for el in container:
            kind = _local(el.tag)
            if kind not in ("sp", "pic", "graphicFrame", "grpSp", "cxnSp"):
                continue
            nv = next((c for c in el.iter() if _local(c.tag) == "cNvPr"), None)
            xfrm = next((c for c in el if _local(c.tag) in ("spPr", "grpSpPr", "xfrm")), None)
            if xfrm is not None and _local(xfrm.tag) != "xfrm":
                xfrm = xfrm.find("a:xfrm", NS)
            box = None
            if xfrm is not None and xfrm.find("a:off", NS) is not None:
                off, ext = xfrm.find("a:off", NS), xfrm.find("a:ext", NS)
                x, y = int(off.get("x")), int(off.get("y"))
                w, h = int(ext.get("cx")), int(ext.get("cy"))
                box = xform(x, y, w, h)
            cid = None
            if nv is not None:
                cr = nv.find(".//a16:creationId", NS)
                cid = cr.get("id") if cr is not None else None
            out.append({
                "shape_id": nv.get("id") if nv is not None else None,
                "shape_name": nv.get("name") if nv is not None else None,
                "creation_id": cid,
                "kind": kind,
                "box": box,
                "text": _text(el)[:200],
            })
            if kind == "grpSp" and xfrm is not None and xfrm.find("a:chOff", NS) is not None and box:
                ch_off, ch_ext = xfrm.find("a:chOff", NS), xfrm.find("a:chExt", NS)
                cx, cy = int(ch_off.get("x")), int(ch_off.get("y"))
                cw, chh = max(int(ch_ext.get("cx")), 1), max(int(ch_ext.get("cy")), 1)
                gx, gy, gw, gh = box
                sx, sy = gw / cw, gh / chh

                def child_xform(x, y, w, h, gx=gx, gy=gy, cx=cx, cy=cy, sx=sx, sy=sy):
                    return (round(gx + (x - cx) * sx), round(gy + (y - cy) * sy), round(w * sx), round(h * sy))

                walk(el, child_xform)

    if tree is not None:
        walk(tree, lambda x, y, w, h: (x, y, w, h))
    return out


def _ph_key(el: ET.Element):
    ph = el.find(".//p:nvPr/p:ph", NS)
    if ph is None:
        return None
    return (ph.get("type", "body"), ph.get("idx", "0"))


def _ph_boxes(z: zipfile.ZipFile, part: str) -> dict:
    """Placeholder boxes a slide inherits: layout first, then master. Keyed by (type, idx) and by type."""
    boxes: dict = {}
    chain, cur = [], part
    for want in ("slideLayout", "slideMaster"):
        nxt = next((t for _i, rt, t in _rels(z, cur) if rt.endswith("/" + want)), None)
        if nxt is None:
            break
        chain.append(nxt)
        cur = nxt
    for src in chain:
        for el in ET.fromstring(z.read(src)).iter():
            if _local(el.tag) not in ("sp", "pic", "graphicFrame"):
                continue
            key = _ph_key(el)
            xfrm = el.find(".//a:xfrm", NS)
            if key is None or xfrm is None or xfrm.find("a:off", NS) is None:
                continue
            off, ext = xfrm.find("a:off", NS), xfrm.find("a:ext", NS)
            box = (int(off.get("x")), int(off.get("y")), int(ext.get("cx")), int(ext.get("cy")))
            boxes.setdefault(key, box)
            boxes.setdefault(key[0], box)
    return boxes


def _inherit_boxes(slide_root: ET.Element, shapes: list[dict], inherited: dict) -> None:
    by_id = {nv.get("id"): el for el in slide_root.iter() if _local(el.tag) in ("sp", "pic", "graphicFrame")
             for nv in [next((c for c in el.iter() if _local(c.tag) == "cNvPr"), None)] if nv is not None}
    for s in shapes:
        if s["box"] is None and s["shape_id"] in by_id:
            key = _ph_key(by_id[s["shape_id"]])
            if key:
                s["box"] = inherited.get(key) or inherited.get(key[0])
                s["box_inherited"] = s["box"] is not None


def _dist(box, px, py) -> float:
    x, y, w, h = box
    dx = max(x - px, 0, px - (x + w))
    dy = max(y - py, 0, py - (y + h))
    return (dx * dx + dy * dy) ** 0.5


def resolve_target(cm: ET.Element, text: str, shapes: list[dict]) -> dict:
    """Order: explicit object moniker > quoted text > pin position. A disagreement between an
    anchor and a quote is surfaced, not arbitrated."""
    result = {"method": "unresolved", "confidence": "none", "shape_id": None,
              "shape_name": None, "shape_text": None, "notes": []}

    def pick(shape, method, confidence):
        result.update(method=method, confidence=confidence, shape_id=shape["shape_id"],
                      shape_name=shape["shape_name"], shape_text=shape["text"][:120])

    anchored = None
    for el in cm.iter():
        name = _local(el.tag)
        if name.endswith("Mk") and name not in ("docMk", "sldMk"):
            sid, cid = el.get("id"), el.get("creationId")
            anchored = next((s for s in shapes if (cid and s["creation_id"] == cid)
                             or (sid and s["shape_id"] == sid)), None)
            if anchored is None:
                result["notes"].append(f"anchor {name} id={sid} creationId={cid} matches no shape on this slide")
            break

    quoted = None
    for q in QUOTE_RE.findall(text):
        hits = [s for s in shapes if q.strip() and q.strip() in s["text"]]
        if len(hits) == 1:
            quoted = hits[0]
            break
        if len(hits) > 1:
            result["notes"].append(f"quote 「{q}」 matches {len(hits)} shapes: " + ", ".join(s["shape_name"] or "?" for s in hits))

    if anchored:
        pick(anchored, "anchor", "high")
        if quoted and quoted is not anchored:
            result["notes"].append(f"quote points at {quoted['shape_name']}, anchor at {anchored['shape_name']}: confirm with the user")
            result["confidence"] = "conflict"
        return result
    if quoted:
        pick(quoted, "quote", "high")
        return result

    pos = cm.find("p188:pos", NS)
    if pos is None:
        pos = next((c for c in cm if _local(c.tag) == "pos"), None)
    if pos is None or _local(cm.tag) != "cm" or cm.find("p188:txBody", NS) is None:
        if pos is not None and cm.find("p188:txBody", NS) is None:
            result["notes"].append("legacy comment: pin units unverified, position not used")
        return result
    px, py = int(pos.get("x")), int(pos.get("y"))
    result["pin_pt"] = [round(px / EMU_PER_PT, 1), round(py / EMU_PER_PT, 1)]
    boxed = [s for s in shapes if s["box"] and _dist(s["box"], px, py) == 0]
    if boxed:
        boxed.sort(key=lambda s: s["box"][2] * s["box"][3])
        pick(boxed[0], "position", "medium" if len(boxed) == 1 else "low")
        if len(boxed) > 1:
            result["notes"].append("pin inside " + ", ".join(s["shape_name"] or "?" for s in boxed) + "; smallest chosen")
        return result
    near = sorted((s for s in shapes if s["box"]), key=lambda s: _dist(s["box"], px, py))
    if near and _dist(near[0]["box"], px, py) <= NEAR_TOLERANCE_EMU:
        pick(near[0], "position-near", "low")
    return result


def read_comments(deck: Path) -> dict:
    with zipfile.ZipFile(deck) as z:
        names = set(z.namelist())
        authors: dict[str, str] = {}
        for part in ("ppt/authors.xml", "ppt/commentAuthors.xml"):
            if part in names:
                for a in ET.fromstring(z.read(part)):
                    authors[a.get("id")] = a.get("name")
        pres_rels = {rid: tgt for rid, _t, tgt in _rels(z, "ppt/presentation.xml")}
        pres = ET.fromstring(z.read("ppt/presentation.xml"))
        slide_parts = [(int(s.get("id")), pres_rels[s.get(f"{{{NS['r']}}}id")])
                       for s in pres.find("p:sldIdLst", NS)]
        comments, per_slide = [], []
        for index, (sld_id, part) in enumerate(slide_parts, start=1):
            slide_root = ET.fromstring(z.read(part))
            shapes = _shapes(slide_root)
            _inherit_boxes(slide_root, shapes, _ph_boxes(z, part))
            count = 0
            for _rid, rtype, target in _rels(z, part):
                if rtype not in (REL_MODERN, REL_LEGACY) or target not in names:
                    continue
                for cm in ET.fromstring(z.read(target)):
                    if _local(cm.tag) != "cm":
                        continue
                    body = cm.find("p188:txBody", NS)
                    text = _text(body) if body is not None else (
                        (next((c for c in cm if _local(c.tag) == "text"), None) or ET.Element("x")).text or "")
                    replies = [{"author": authors.get(r.get("authorId"), r.get("authorId")),
                                "text": _text(r.find("p188:txBody", NS)) if r.find("p188:txBody", NS) is not None else ""}
                               for r in cm.iter(f"{{{NS['p188']}}}reply")]
                    status = cm.get("status", "active")
                    entry = {
                        "id": cm.get("id") or cm.get("idx"),
                        "slide": index,
                        "slide_id": sld_id,
                        "author": authors.get(cm.get("authorId"), cm.get("authorId")),
                        "created": cm.get("created") or cm.get("dt"),
                        "status": status,
                        "format": "modern" if rtype == REL_MODERN else "legacy",
                        "text": text,
                        "replies": replies,
                        "target": resolve_target(cm, text, shapes) if status != "resolved" else None,
                    }
                    comments.append(entry)
                    count += 1
            per_slide.append({"slide": index, "comments": count})
    open_ = [c for c in comments if c["status"] != "resolved"]
    return {
        "deck": str(deck.resolve()),
        "sha256": hashlib.sha256(deck.read_bytes()).hexdigest(),
        "summary": {
            "total": len(comments),
            "open": len(open_),
            "resolved_in_powerpoint": len(comments) - len(open_),
            "by_method": {m: sum(1 for c in open_ if c["target"]["method"] == m)
                          for m in ("anchor", "quote", "position", "position-near", "unresolved")},
        },
        "comments": comments,
    }


# ---------------------------------------------------------------- cursor

SELECTION = {0: "none", 1: "slides", 2: "shapes", 3: "text"}
VIEW_SLIDESHOW = "slideshow"


def read_cursor(app) -> dict:
    """`app` is a PowerPoint.Application COM object (or a test double with the same surface)."""
    state: dict = {"state": "ok"}
    if app.Presentations.Count == 0:
        return {"state": "no-presentation"}
    if app.SlideShowWindows.Count > 0:
        show = app.SlideShowWindows(1)
        pres = show.Presentation
        state.update(view=VIEW_SLIDESHOW, slide=show.View.CurrentShowPosition, selection="none", shapes=[])
    else:
        win = app.ActiveWindow
        pres = win.Presentation
        state["view"] = "normal"
        try:
            state["slide"] = win.View.Slide.SlideIndex
        except Exception:
            state["slide"] = None
            state["notes_"] = ["active pane holds no slide (outline/notes/sorter?)"]
        sel = win.Selection
        state["selection"] = SELECTION.get(sel.Type, str(sel.Type))
        shapes = []
        if sel.Type in (2, 3):
            rng = sel.ShapeRange
            for i in range(1, rng.Count + 1):
                shp = rng.Item(i)
                txt = ""
                try:
                    if shp.HasTextFrame:
                        txt = shp.TextFrame.TextRange.Text
                except Exception:
                    pass
                shapes.append({"shape_id": str(shp.Id), "shape_name": shp.Name, "text": txt[:120].replace("\r", "\n")})
        state["shapes"] = shapes
        if sel.Type == 3:
            state["selected_text"] = sel.TextRange.Text[:200].replace("\r", "\n")
        if sel.Type == 1:
            state["slides_selected"] = [sel.SlideRange.Item(i).SlideIndex for i in range(1, sel.SlideRange.Count + 1)]
    state.update(deck=pres.FullName, total_slides=pres.Slides.Count, saved=bool(pres.Saved))
    if not state["saved"]:
        state.setdefault("notes_", []).append("unsaved changes: comments on disk may lag what is on screen")
    return state


def attach_powerpoint():
    import pythoncom
    import win32com.client
    pythoncom.CoInitialize()
    try:
        return win32com.client.GetActiveObject("PowerPoint.Application")
    except Exception:
        return None


# ---------------------------------------------------------------- selftest

def _build_fixture(path: Path) -> None:
    from pptx import Presentation
    from pptx.util import Pt

    prs = Presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    t = slide.shapes.add_textbox(Pt(100), Pt(100), Pt(300), Pt(60))
    t.name, t.text_frame.text = "title_main", "Probe title"
    b = slide.shapes.add_textbox(Pt(100), Pt(300), Pt(300), Pt(60))
    b.name, b.text_frame.text = "body_1", "Body text here"
    d = slide.shapes.add_textbox(Pt(500), Pt(300), Pt(100), Pt(40))
    d.name, d.text_frame.text = "dup_a", "Body text here too"
    s2 = prs.slides.add_slide(prs.slide_layouts[5])  # "Title Only": title placeholder, no xfrm on the slide
    s2.shapes.title.text = "Inherited title"
    prs.save(path)

    ids = {}
    with zipfile.ZipFile(path) as z:
        items = {n: z.read(n) for n in z.namelist()}
    slide_xml = items["ppt/slides/slide1.xml"].decode()
    for m in re.finditer(r'<p:cNvPr id="(\d+)" name="([^"]+)"', slide_xml):
        ids[m.group(2)] = m.group(1)
    # give body_1 a creationId the way PowerPoint does on save
    slide_xml = slide_xml.replace(
        f'<p:cNvPr id="{ids["body_1"]}" name="body_1"/>',
        f'<p:cNvPr id="{ids["body_1"]}" name="body_1"><a:extLst><a:ext uri="{{FF2B5EF4-FFF2-40B4-BE49-F238E27FC236}}">'
        f'<a16:creationId xmlns:a16="{NS["a16"]}" id="{{BODY-CID}}"/></a:ext></a:extLst></p:cNvPr>')
    items["ppt/slides/slide1.xml"] = slide_xml.encode()

    def cm(cid, text, pos=None, anchor=None, status=None):
        mk = f'<ac:spMk xmlns:ac="http://schemas.microsoft.com/office/drawing/2013/main/command" id="{anchor[0]}" creationId="{anchor[1]}"/>' if anchor else ""
        st = f' status="{status}"' if status else ""
        p = f'<p188:pos x="{pos[0] * EMU_PER_PT}" y="{pos[1] * EMU_PER_PT}"/>' if pos else ""
        return (f'<p188:cm id="{{{cid}}}" authorId="{{AU}}" created="2026-09-29T00:00:00"{st}>'
                f'<pc:sldMkLst xmlns:pc="http://schemas.microsoft.com/office/powerpoint/2013/main/command"><pc:docMk/><pc:sldMk cId="1" sldId="256"/></pc:sldMkLst>'
                f'{mk}{p}<p188:txBody><a:bodyPr/><a:lstStyle/><a:p><a:r><a:t>{text}</a:t></a:r></a:p></p188:txBody></p188:cm>')

    body = "".join([
        cm("C1", "make this red", pos=(110, 110)),                                  # inside title_main
        cm("C2", "shorten", anchor=("999", "{BODY-CID}")),                          # anchor via creationId
        cm("C3", "what is this", pos=(900, 900)),                                   # nowhere
        cm("C4", "「Probe title」 should be bigger", pos=(900, 900)),               # quote wins over empty pin
        cm("C5", "done already", pos=(110, 110), status="resolved"),                # skipped
        cm("C6", "tighten", pos=(410, 110)),                                        # 10pt right of title: near
        cm("C7", "fix 「Body text here」", pos=(900, 900)),                         # ambiguous quote: 2 shapes
        cm("C8", "same as quote?", anchor=("999", "{BODY-CID}")).replace("same as quote?", "「Probe title」 vs anchor"),
    ])
    items["ppt/comments/modernComment_100_T.xml"] = (
        f'<?xml version="1.0" encoding="UTF-8" standalone="yes"?><p188:cmLst xmlns:a="{NS["a"]}" '
        f'xmlns:p188="{NS["p188"]}">{body}</p188:cmLst>').encode()
    items["ppt/authors.xml"] = (f'<?xml version="1.0" encoding="UTF-8"?><p188:authorLst xmlns:p188="{NS["p188"]}">'
                                f'<p188:author id="{{AU}}" name="Reviewer" initials="R" userId="r" providerId="None"/></p188:authorLst>').encode()
    items["ppt/comments/modernComment_101_T.xml"] = (
        f'<?xml version="1.0" encoding="UTF-8" standalone="yes"?><p188:cmLst xmlns:a="{NS["a"]}" '
        f'xmlns:p188="{NS["p188"]}">{cm("C9", "bigger", pos=(300, 50))}</p188:cmLst>').encode()
    for n, part in ((1, "modernComment_100_T.xml"), (2, "modernComment_101_T.xml")):
        rels = items[f"ppt/slides/_rels/slide{n}.xml.rels"].decode()
        items[f"ppt/slides/_rels/slide{n}.xml.rels"] = rels.replace(
            "</Relationships>", f'<Relationship Id="rIdCm" Type="{REL_MODERN}" Target="../comments/{part}"/></Relationships>').encode()
    ct = items["[Content_Types].xml"].decode()
    items["[Content_Types].xml"] = ct.replace("</Types>", (
        '<Override PartName="/ppt/comments/modernComment_100_T.xml" ContentType="application/vnd.ms-powerpoint.comments+xml"/>'
        '<Override PartName="/ppt/comments/modernComment_101_T.xml" ContentType="application/vnd.ms-powerpoint.comments+xml"/>'
        '<Override PartName="/ppt/authors.xml" ContentType="application/vnd.ms-powerpoint.authors+xml"/></Types>')).encode()
    with zipfile.ZipFile(path, "w", zipfile.ZIP_DEFLATED) as z:
        for n, data in items.items():
            z.writestr(n, data)


class _Fake:
    def __init__(self, **kw):
        self.__dict__.update(kw)


class _Coll:
    def __init__(self, items):
        self._items = list(items)
        self.Count = len(self._items)

    def __call__(self, i):
        return self._items[i - 1]

    Item = __call__


def _fake_app(sel_type: int, slideshow: bool = False, saved: bool = True):
    pres = _Fake(FullName=r"D:\x\deck.pptx", Slides=_Coll([1, 2, 3]), Saved=saved)
    shp = _Fake(Id=7, Name="title_main", HasTextFrame=True, TextFrame=_Fake(TextRange=_Fake(Text="Probe\rtitle")))
    sel = _Fake(Type=sel_type, ShapeRange=_Coll([shp]), TextRange=_Fake(Text="Pro"),
                SlideRange=_Coll([_Fake(SlideIndex=2)]))
    win = _Fake(Presentation=pres, View=_Fake(Slide=_Fake(SlideIndex=2)), Selection=sel)
    shows = _Coll([_Fake(Presentation=pres, View=_Fake(CurrentShowPosition=3))] if slideshow else [])
    return _Fake(Presentations=_Coll([pres]), SlideShowWindows=shows, ActiveWindow=win)


def selftest() -> int:
    import tempfile

    failures = []

    def check(label, got, want):
        ok = got == want
        print(f"  {'PASS' if ok else 'FAIL'}  {label}: got {got!r}" + ("" if ok else f", want {want!r}"))
        if not ok:
            failures.append(label)

    with tempfile.TemporaryDirectory() as tmp:
        fx = Path(tmp) / "fixture.pptx"
        _build_fixture(fx)
        before = fx.read_bytes()
        res = read_comments(fx)
        by = {c["id"].strip("{}"): c for c in res["comments"]}
        print("comments:")
        check("C1 pin inside title -> position/title_main", (by["C1"]["target"]["method"], by["C1"]["target"]["shape_name"]), ("position", "title_main"))
        check("C2 creationId anchor -> anchor/body_1", (by["C2"]["target"]["method"], by["C2"]["target"]["shape_name"]), ("anchor", "body_1"))
        check("C3 pin in empty space -> unresolved (negative control)", by["C3"]["target"]["method"], "unresolved")
        check("C4 quote beats empty pin -> quote/title_main", (by["C4"]["target"]["method"], by["C4"]["target"]["shape_name"]), ("quote", "title_main"))
        check("C5 resolved in PowerPoint -> no target", (by["C5"]["status"], by["C5"]["target"]), ("resolved", None))
        check("C6 pin 10pt outside -> position-near/low", (by["C6"]["target"]["method"], by["C6"]["target"]["confidence"]), ("position-near", "low"))
        check("C7 quote on 2 shapes -> unresolved + note", (by["C7"]["target"]["method"], bool(by["C7"]["target"]["notes"])), ("unresolved", True))
        check("C8 anchor vs quote disagree -> conflict", by["C8"]["target"]["confidence"], "conflict")
        t9 = by["C9"]["target"]
        check("C9 pin on inherited-title placeholder -> position via layout box",
              (by["C9"]["slide"], t9["method"], t9["shape_text"]), (2, "position", "Inherited title"))
        check("summary open count", res["summary"]["open"], 8)
        check("author resolved from authors.xml", by["C1"]["author"], "Reviewer")
        check("deck bytes untouched", fx.read_bytes() == before, True)

    print("cursor:")
    check("no presentation", read_cursor(_Fake(Presentations=_Coll([]), SlideShowWindows=_Coll([])))["state"], "no-presentation")
    c = read_cursor(_fake_app(2))
    check("shape selection", (c["slide"], c["selection"], c["shapes"][0]["shape_name"], c["shapes"][0]["text"]), (2, "shapes", "title_main", "Probe\ntitle"))
    c = read_cursor(_fake_app(3))
    check("text selection carries selected_text", (c["selection"], c.get("selected_text")), ("text", "Pro"))
    c = read_cursor(_fake_app(0, saved=False))
    check("unsaved deck is flagged", (c["saved"], any("unsaved" in n for n in c.get("notes_", []))), (False, True))
    c = read_cursor(_fake_app(2, slideshow=True))
    check("slideshow uses show position", (c["view"], c["slide"]), ("slideshow", 3))
    c = read_cursor(_fake_app(1))
    check("slide-sorter selection lists slides", c.get("slides_selected"), [2])

    print(f"selftest: {'OK' if not failures else 'FAILED ' + ', '.join(failures)}")
    return 0 if not failures else 1


def _emit(data: dict, out: str | None) -> None:
    text = json.dumps(data, ensure_ascii=False, indent=2)
    if out:
        Path(out).write_text(text, encoding="utf-8")
        print(out)
    else:
        print(text)


def main() -> int:
    ap = argparse.ArgumentParser(description=__doc__.split("\n")[0])
    sub = ap.add_subparsers(dest="cmd", required=True)
    c = sub.add_parser("comments", help="read and resolve reviewer comments in a .pptx")
    c.add_argument("deck")
    c.add_argument("--out")
    k = sub.add_parser("cursor", help="where the user is in the running PowerPoint")
    k.add_argument("--out")
    sub.add_parser("selftest")
    args = ap.parse_args()

    if args.cmd == "selftest":
        return selftest()
    if args.cmd == "comments":
        deck = Path(args.deck)
        if not deck.is_file():
            print(f"pptx-review: no such deck: {deck}", file=sys.stderr)
            return 2
        _emit(read_comments(deck), args.out)
        return 0
    app = attach_powerpoint()
    if app is None:
        _emit({"state": "no-powerpoint", "hint": "PowerPoint is not running; open the deck first"}, args.out)
        return 3
    _emit(read_cursor(app), args.out)
    return 0


if __name__ == "__main__":
    sys.exit(main())
